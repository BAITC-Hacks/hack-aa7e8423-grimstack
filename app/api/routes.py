"""Procurement API. Business calculations stay in app.engine."""

import math
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from starlette.datastructures import UploadFile

from app import ai, engine, ingest
from app.api.schemas import ApprovalRecord, ChangedLine, CompareDelta, CompareRequest, CompareResult
from app.engine.config import SUPPLIERS
from app.ingest.dataset import rename_supplier
from app.contracts import (
    DatasetUploaded,
    ErrorBody,
    LinePatch,
    Meta,
    OrderLine,
    RunParams,
    RunResult,
    SkuHistory,
    SummaryResponse,
    Supplier,
    SupplierSummary,
)
from app.store import Store


router = APIRouter()
BACKTEST_REPORT_PATH = Path(__file__).resolve().parents[2] / "data" / "backtest" / "report.json"
MAX_UPLOAD_BYTES = 30 * 1024 * 1024
EXPORT_COLUMNS = (
    "Код 1С", "Артикул поставщика", "Наименование", "Ед.", "Количество",
    "Цена", "Сумма", "Поставщик", "Срочность", "Обоснование",
)
ERROR_RESPONSES = {
    404: {"model": ErrorBody, "description": "Ресурс не найден"},
    409: {"model": ErrorBody, "description": "Заказ уже утверждён"},
    422: {"model": ErrorBody, "description": "Некорректные данные запроса"},
    503: {"model": ErrorBody, "description": "Сводка ИИ недоступна"},
}
UPLOAD_REQUEST_BODY = {
    "requestBody": {
        "required": True,
        "content": {
            "multipart/form-data": {
                "schema": {
                    "type": "object",
                    "properties": {
                        "monthly_sales": {"type": "string", "format": "binary", "description": "Продажи по месяцам"},
                        "monthly_stock": {"type": "string", "format": "binary", "description": "Остатки по месяцам"},
                        "in_transit": {"type": "string", "format": "binary", "description": "Товары в пути"},
                        "moq": {"type": "string", "format": "binary", "description": "Кратность заказа"},
                        "sales_tx": {"type": "string", "format": "binary", "description": "Строки продаж, если есть"},
                        "seasonality": {"type": "string", "format": "binary", "description": "Сезонность, если есть"},
                    },
                    "required": ["monthly_sales", "monthly_stock", "in_transit", "moq"],
                }
            }
        },
    }
}
NEW_SUPPLIER_UPLOAD_BODY = {
    "requestBody": {
        "required": True,
        "content": {
            "multipart/form-data": {
                "schema": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string", "description": "Название нового поставщика"},
                        "template": {"type": "string", "enum": ["IEK", "SE"],
                                     "description": "Формат выгрузки 1С"},
                        **UPLOAD_REQUEST_BODY["requestBody"]["content"]["multipart/form-data"]["schema"]["properties"],
                    },
                    "required": ["name", "template", "monthly_sales", "monthly_stock", "in_transit", "moq"],
                }
            }
        },
    }
}
CUSTOM_SUPPLIER_CODE = re.compile(r"CUSTOM_[0-9a-f]{8}\Z")


def errors(*statuses: int) -> dict[int, dict]:
    return {status: ERROR_RESPONSES[status] for status in statuses}


def fit_columns(sheet) -> None:
    for column in sheet.columns:
        longest = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(80, max(12, longest + 2))


def get_store(request: Request) -> Store:
    return request.app.state.store


def parse_supplier(raw: str) -> Supplier:
    # Path values are checked here instead of Literal validation: unknown paths are 404.
    if raw not in SUPPLIERS and not CUSTOM_SUPPLIER_CODE.fullmatch(raw):
        raise HTTPException(status_code=404, detail="Поставщик не найден")
    return raw


def get_run(store: Store, run_id: str) -> RunResult:
    try:
        return store.runs[run_id]
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Расчёт не найден") from exc


def get_supplier(run: RunResult, supplier: Supplier) -> SupplierSummary:
    found = next((s for s in run.suppliers if s.supplier == supplier), None)
    if found is None:
        raise HTTPException(status_code=404, detail="Поставщик не найден в расчёте")
    return found


def validate_run_supplier(dataset, supplier: str | None) -> None:
    if supplier is None:
        return
    skus = getattr(dataset, "skus", None)
    available = (
        set(skus.index.get_level_values("supplier"))
        if skus is not None else set(SUPPLIERS)
    )
    if supplier not in available:
        raise HTTPException(status_code=422, detail="Поставщик отсутствует в выбранном наборе данных")


async def read_upload_files(request: Request, *, new_supplier: bool = False) -> tuple[dict[str, bytes], dict[str, str]]:
    form = await request.form()
    files: dict[str, bytes] = {}
    fields: dict[str, str] = {}
    for role, item in form.multi_items():
        if role in files or role in fields:
            raise HTTPException(status_code=422, detail=f"Поле {role} передано дважды")
        if new_supplier and role in ("name", "template"):
            if not isinstance(item, str):
                raise HTTPException(status_code=422, detail=f"Поле {role} должно быть текстом")
            fields[role] = item
            continue
        if not isinstance(item, UploadFile):
            raise HTTPException(status_code=422, detail=f"Поле {role} должно быть файлом")
        content = await item.read(MAX_UPLOAD_BYTES + 1)
        if len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=422, detail=f"Файл {role} больше 30 МБ")
        files[role] = content
    return files, fields


def upload_warnings(dataset, files: dict[str, bytes]) -> list[str]:
    warnings = dataset.get("warnings", []) if isinstance(dataset, dict) else getattr(dataset, "warnings", [])
    if "sales_tx" not in files:
        warnings = [*warnings, "Без строк продаж очистка разовых заказов и оценка дней наличия ограничены"]
    return list(warnings)


@router.get(
    "/meta", response_model=Meta,
    tags=["Данные"], summary="Получить параметры исходных данных",
    description="Показывает дату данных и поставщиков выбранного набора; без dataset_id использует встроенные данные.",
    responses=errors(404),
)
def meta(dataset_id: str | None = None, store: Store = Depends(get_store)) -> Meta:
    try:
        dataset = store.dataset_for(dataset_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Набор данных не найден") from exc
    return engine.meta(dataset)


@router.get(
    "/backtest", tags=["Бэктест"], summary="Получить отчёт бэктеста",
    description="Отдаёт готовый JSON-отчёт проверки расчёта на исторических данных без пересчёта.",
    responses={
        200: {"description": "Исходный JSON-отчёт", "content": {"application/json": {"schema": {"type": "object"}}}},
        **errors(404),
    },
)
def backtest_report() -> Response:
    if not BACKTEST_REPORT_PATH.is_file():
        error = ErrorBody(detail="Отчёт бэктеста не найден", code="backtest_not_found")
        return JSONResponse(status_code=404, content=error.model_dump(mode="json"))
    return FileResponse(BACKTEST_REPORT_PATH, media_type="application/json")


@router.post(
    "/runs", response_model=RunResult,
    tags=["Расчёты"], summary="Рассчитать заказ поставщикам",
    description="Создаёт и сохраняет расчёт по выбранному набору данных и параметрам.",
    responses=errors(404, 422),
)
def create_run(params: RunParams, store: Store = Depends(get_store)) -> RunResult:
    try:
        dataset = store.dataset_for(params.dataset_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Набор данных не найден") from exc
    validate_run_supplier(dataset, params.supplier)
    result = engine.run(dataset, params)
    store.save_run(result, dataset)
    return result


@router.post(
    "/runs/compare", response_model=CompareResult,
    tags=["Расчёты"], summary="Сравнить два сценария расчёта",
    description="Считает и сохраняет базовый и изменённый сценарии, возвращает разницу и товары с наибольшим изменением заказа.",
    responses=errors(404, 422),
)
def compare_runs(request: CompareRequest, store: Store = Depends(get_store)) -> CompareResult:
    try:
        base_dataset = store.dataset_for(request.base.dataset_id)
        scenario_dataset = store.dataset_for(request.scenario.dataset_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Набор данных не найден") from exc

    validate_run_supplier(base_dataset, request.base.supplier)
    validate_run_supplier(scenario_dataset, request.scenario.supplier)

    base = engine.run(base_dataset, request.base)
    scenario = engine.run(scenario_dataset, request.scenario)
    with store.lock:
        store.save_run(base, base_dataset)
        store.save_run(scenario, scenario_dataset)

    base_lines = {line.line_id: line for line in base.lines}
    scenario_lines = {line.line_id: line for line in scenario.lines}
    changed = []
    for line_id in base_lines.keys() | scenario_lines.keys():
        before = base_lines.get(line_id)
        after = scenario_lines.get(line_id)
        base_qty = before.recommended_qty if before else 0.0
        scenario_qty = after.recommended_qty if after else 0.0
        base_urgency = before.urgency if before else "none"
        scenario_urgency = after.urgency if after else "none"
        if base_qty == scenario_qty and base_urgency == scenario_urgency:
            continue
        line = after or before
        assert line is not None
        changed.append(ChangedLine(
            line_id=line_id,
            name=line.name,
            supplier=line.supplier,
            base_qty=base_qty,
            scenario_qty=scenario_qty,
            base_urgency=base_urgency,
            scenario_urgency=scenario_urgency,
        ))
    changed.sort(key=lambda line: (-abs(line.scenario_qty - line.base_qty), line.line_id))

    amount = (
        round(scenario.kpi.total_amount - base.kpi.total_amount, 2)
        if base.kpi.total_amount is not None and scenario.kpi.total_amount is not None
        else None
    )
    return CompareResult(
        base_run_id=base.run_id,
        scenario_run_id=scenario.run_id,
        delta=CompareDelta(
            lines_to_order=scenario.kpi.lines_to_order - base.kpi.lines_to_order,
            critical=scenario.kpi.critical - base.kpi.critical,
            total_qty=(
                sum(item.total_qty for item in scenario.suppliers)
                - sum(item.total_qty for item in base.suppliers)
            ),
            total_amount=amount,
        ),
        changed=changed[:20],
    )


@router.get(
    "/runs/{run_id}", response_model=RunResult,
    tags=["Расчёты"], summary="Получить сохранённый расчёт",
    description="Возвращает строки заказа, итоги и параметры ранее созданного расчёта.",
    responses=errors(404),
)
def read_run(run_id: str, store: Store = Depends(get_store)) -> RunResult:
    with store.lock:
        return get_run(store, run_id).model_copy(deep=True)


@router.patch(
    "/runs/{run_id}/lines/{line_id}", response_model=OrderLine,
    tags=["Заказы"], summary="Изменить количество товара в заказе",
    description="Проверяет верхний предел и кратность упаковки, затем обновляет количество и итоги неутверждённого заказа.",
    responses=errors(404, 409, 422),
)
def patch_line(
    run_id: str, line_id: str, patch: LinePatch, store: Store = Depends(get_store)
) -> OrderLine:
    with store.lock:
        run = get_run(store, run_id)
        line = next((item for item in run.lines if item.line_id == line_id), None)
        if line is None:
            raise HTTPException(status_code=404, detail="Строка заказа не найдена")
        if get_supplier(run, line.supplier).status == "approved":
            raise HTTPException(status_code=409, detail="Заказ поставщика уже утверждён")

        quantity = patch.final_qty
        multiple = line.moq
        if not math.isfinite(quantity) or multiple <= 0 or not math.isfinite(multiple):
            raise HTTPException(status_code=422, detail="Некорректное количество или кратность")
        quantity_limit = max(100 * line.recommended_qty, 100 * multiple, 1000)
        if quantity > quantity_limit:
            raise HTTPException(
                status_code=422,
                detail=f"Количество не должно превышать {quantity_limit:g} для этого товара",
            )
        rounded = round(quantity / multiple) * multiple
        if not math.isclose(quantity, rounded, rel_tol=1e-9):
            raise HTTPException(
                status_code=422, detail=f"Количество должно быть кратно {multiple:g}"
            )

        line.final_qty = quantity
        line.amount = (
            round(quantity * line.unit_cost, 2)
            if line.unit_cost is not None and math.isfinite(line.unit_cost)
            else None
        )
        store.refresh_totals(run)
        return line.model_copy(deep=True)


@router.post(
    "/runs/{run_id}/suppliers/{supplier}/approve", response_model=SupplierSummary,
    tags=["Заказы"], summary="Утвердить заказ поставщику",
    description="Сохраняет утверждение и состав заказа; повторное утверждение возвращает прежний результат.",
    responses=errors(404),
)
def approve_supplier(
    run_id: str, supplier: str, store: Store = Depends(get_store)
) -> SupplierSummary:
    selected = parse_supplier(supplier)
    with store.lock:
        run = get_run(store, run_id)
        summary = get_supplier(run, selected)
        if summary.status == "approved":
            return summary.model_copy(deep=True)

        approved_at = datetime.now(timezone.utc)
        lines = [
            line.model_dump(mode="json")
            for line in run.lines
            if line.supplier == selected and line.final_qty > 0
        ]
        saved = store.append_approval({
            "run_id": run_id,
            "supplier": selected,
            "approved_at": approved_at.isoformat(),
            "data_as_of": run.data_as_of.isoformat(),
            "lines": lines,
        })
        summary.status = "approved"
        summary.approved_at = datetime.fromisoformat(saved["approved_at"])
        return summary.model_copy(deep=True)


@router.get(
    "/approvals", response_model=list[ApprovalRecord],
    tags=["Заказы"], summary="Получить историю утверждённых заказов",
    description="Возвращает сохранённые в SQLite утверждения и состав заказов, начиная с последних.",
)
def approval_history(store: Store = Depends(get_store)) -> list[dict]:
    return store.list_approvals()


@router.get(
    "/runs/{run_id}/export.xlsx",
    tags=["Заказы"], summary="Скачать заказ в Excel для 1С",
    description="Выгружает положительные строки выбранного поставщика без отправки заказа поставщику.",
    responses={
        200: {"description": "Файл Excel", "content": {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {
                "schema": {"type": "string", "format": "binary"}
            }
        }},
        **errors(404, 422),
    },
)
def export_run(
    run_id: str, supplier: str, store: Store = Depends(get_store)
) -> Response:
    selected = parse_supplier(supplier)
    with store.lock:
        run = get_run(store, run_id)
        summary = get_supplier(run, selected).model_copy(deep=True)
        params = run.params.model_copy(deep=True)
        data_as_of = run.data_as_of
        dataset = store.run_datasets[run_id]
        custom = getattr(dataset, "metadata", {}).get("custom_suppliers", {}).get(selected)
        defaults = custom or SUPPLIERS.get(selected, {})
        lines = [
            line.model_copy(deep=True)
            for line in run.lines
            if line.supplier == selected and line.final_qty > 0
        ]
        filename = f"order_{selected}_{run.created_at.date().isoformat()}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = selected
    sheet.append(EXPORT_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for line in lines:
        values = (
            line.sku, line.article, line.name, line.unit, line.final_qty,
            line.unit_cost, line.amount, line.supplier, line.urgency, line.explanation,
        )
        sheet.append(values)
        # Prevent an uploaded SKU/name/explanation starting with '=' from becoming a formula.
        for index in (1, 2, 3, 4, 8, 9, 10):
            cell = sheet.cell(sheet.max_row, index)
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for index, number_format in ((5, "#,##0.###"), (6, "#,##0.00"), (7, "#,##0.00")):
            sheet.cell(sheet.max_row, index).number_format = number_format
    last_data_row = sheet.max_row
    sheet.append(("Итого", None, None, None, summary.total_qty, None, summary.total_amount))
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    sheet.cell(sheet.max_row, 5).number_format = "#,##0.###"
    sheet.cell(sheet.max_row, 7).number_format = "#,##0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{last_data_row}"
    fit_columns(sheet)

    settings = workbook.create_sheet("Параметры расчёта")
    settings.append(("Параметр", "Значение"))
    settings_rows = (
        ("Дата данных", data_as_of.isoformat()),
        ("Метод", params.method),
        ("L, дни", params.lead_time_days or defaults.get("lead_time_days")),
        ("R, дни", params.review_period_days or defaults.get("review_period_days")),
        ("Поставщик", summary.supplier_name if custom else selected),
        ("Статус", summary.status),
        ("Дата утверждения", summary.approved_at.isoformat() if summary.approved_at else None),
    )
    for row in settings_rows:
        settings.append(row)
        value = settings.cell(settings.max_row, 2)
        if isinstance(value.value, str):
            value.data_type = "s"
    for cell in settings[1]:
        cell.font = Font(bold=True)
    settings.freeze_panes = "A2"
    fit_columns(settings)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/sku/{supplier}/{sku}/history", response_model=SkuHistory,
    tags=["История"], summary="Получить историю товара",
    description="Возвращает продажи, восстановленный спрос и прогноз выбранного товара.",
    responses=errors(404),
)
def sku_history(
    supplier: str, sku: str, run_id: str | None = None, store: Store = Depends(get_store)
) -> SkuHistory:
    selected = parse_supplier(supplier)
    with store.lock:
        if run_id is None:
            dataset, params = store.default_dataset, RunParams()
        else:
            run = get_run(store, run_id)
            get_supplier(run, selected)
            dataset, params = store.run_datasets[run_id], run.params
    try:
        return engine.history(dataset, selected, sku, params)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Товар не найден") from exc


@router.post(
    "/datasets/new", response_model=DatasetUploaded,
    tags=["Данные"], summary="Добавить нового поставщика из выгрузок 1С",
    description="Создаёт отдельного поставщика по файлам в формате IEK или Systeme Electric; файлы до 30 МБ каждый.",
    responses=errors(422), openapi_extra=NEW_SUPPLIER_UPLOAD_BODY,
)
async def upload_new_supplier(
    request: Request, store: Store = Depends(get_store)
) -> DatasetUploaded:
    files, fields = await read_upload_files(request, new_supplier=True)
    name = fields.get("name", "").strip()
    template = fields.get("template", "")
    if len(name) < 2 or len(name) > 80 or any(ord(char) < 32 for char in name):
        raise HTTPException(status_code=422, detail="Укажите название поставщика длиной от 2 до 80 символов")
    if template not in SUPPLIERS:
        raise HTTPException(status_code=422, detail="Выберите формат выгрузки IEK или Systeme Electric")

    dataset = ingest.load_uploaded(template, files)
    with store.lock:
        supplier_code = f"CUSTOM_{uuid4().hex[:8]}"
        while any(
            supplier_code in getattr(item, "metadata", {}).get("custom_suppliers", {})
            for item in store.datasets.values()
        ):
            supplier_code = f"CUSTOM_{uuid4().hex[:8]}"
        dataset_id = uuid4().hex[:12]
        while dataset_id in store.datasets:
            dataset_id = uuid4().hex[:12]
        config = SUPPLIERS[template]
        rename_supplier(
            dataset, supplier_code, name=name, template=template,
            lead_time_days=config["lead_time_days"],
            review_period_days=config["review_period_days"],
        )
        store.datasets[dataset_id] = dataset
    return DatasetUploaded(
        dataset_id=dataset_id, supplier=supplier_code, supplier_name=name,
        warnings=upload_warnings(dataset, files),
    )


@router.post(
    "/datasets/{supplier}", response_model=DatasetUploaded,
    tags=["Данные"], summary="Загрузить выгрузки поставщика",
    description="Принимает файлы Excel по ролям для IEK или Systeme Electric, не более 30 МБ на файл.",
    responses=errors(404, 422), openapi_extra=UPLOAD_REQUEST_BODY,
)
async def upload_dataset(
    supplier: str, request: Request, store: Store = Depends(get_store)
) -> DatasetUploaded:
    selected = parse_supplier(supplier)
    if selected not in SUPPLIERS:
        raise HTTPException(status_code=404, detail="Для нового поставщика используйте /api/datasets/new")
    files, _ = await read_upload_files(request)

    dataset = ingest.load_uploaded(selected, files)
    dataset_id = uuid4().hex[:12]
    with store.lock:
        store.datasets[dataset_id] = dataset
    return DatasetUploaded(
        dataset_id=dataset_id, supplier=selected, supplier_name=SUPPLIERS[selected]["name"],
        warnings=upload_warnings(dataset, files),
    )


@router.post(
    "/runs/{run_id}/summary", response_model=SummaryResponse,
    tags=["ИИ"], summary="Составить текстовую сводку расчёта",
    description="Возвращает сохранённое или созданное ИИ объяснение заказа выбранного поставщика.",
    responses=errors(404, 422, 503),
)
def summarize_run(
    run_id: str, supplier: str, store: Store = Depends(get_store)
) -> SummaryResponse:
    selected = parse_supplier(supplier)
    with store.lock:
        run = get_run(store, run_id)
        get_supplier(run, selected)
        snapshot = run.model_copy(deep=True)
    summary = ai.summary(snapshot, selected)
    if summary is None:
        raise HTTPException(status_code=503, detail="LLM недоступна: нет ключа и кэша")
    return summary
