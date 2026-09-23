from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.store import store


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def run(client):
    response = client.post("/api/runs", json={"supplier": "SE"})
    assert response.status_code == 200
    return response.json()


def test_health_meta_and_spa(client):
    assert client.get("/health").json() == {"status": "ok"}
    assert {item["supplier"] for item in client.get("/api/meta").json()["suppliers"]} == {"IEK", "SE"}
    assert client.get("/orders").status_code == 200
    assert client.get("/api/not-found").json()["code"] == "not_found"


def test_run_history_and_missing_ids(client, run):
    run_id = run["run_id"]
    assert client.get(f"/api/runs/{run_id}").json()["run_id"] == run_id
    line = run["lines"][0]
    history = client.get(f"/api/sku/SE/{line['sku']}/history")
    assert history.status_code == 200
    assert history.json()["sku"] == line["sku"]
    assert client.get("/api/sku/SE/does-not-exist/history").status_code == 404
    assert client.get("/api/runs/missing").status_code == 404
    assert client.post("/api/runs", json={"dataset_id": "missing"}).status_code == 404


def test_patch_recalculates_totals_and_rejects_invalid(client, run):
    run_id = run["run_id"]
    line = next(item for item in run["lines"] if item["unit_cost"] is not None)
    path = f"/api/runs/{run_id}/lines/{line['line_id']}"
    assert client.patch(path, json={"final_qty": -1}).status_code == 422
    assert client.patch(path, json={"final_qty": line["moq"] / 2}).status_code == 422
    new_qty = line["moq"]
    assert client.patch(path, json={"final_qty": new_qty}).json()["final_qty"] == new_qty
    updated = client.get(f"/api/runs/{run_id}").json()
    assert updated["kpi"]["total_amount"] == updated["suppliers"][0]["total_amount"]
    assert updated["suppliers"][0]["total_qty"] == sum(item["final_qty"] for item in updated["lines"])
    assert client.patch(f"/api/runs/{run_id}/lines/missing", json={"final_qty": 0}).status_code == 404


def test_export_approve_and_conflict(client, run, tmp_path, monkeypatch):
    import app.store as store_module
    monkeypatch.setattr(store_module, "APPROVALS", tmp_path / "approvals.json")
    run_id = run["run_id"]
    response = client.get(f"/api/runs/{run_id}/export.xlsx", params={"supplier": "SE"})
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet["A1"].value == "Код 1С"
    assert sheet.max_row == 1 + sum(line["final_qty"] > 0 for line in run["lines"])
    assert client.post(f"/api/runs/{run_id}/suppliers/SE/approve").json()["status"] == "approved"
    assert (tmp_path / "approvals.json").exists()
    line = run["lines"][0]
    assert client.patch(f"/api/runs/{run_id}/lines/{line['line_id']}", json={"final_qty": 0}).status_code == 409
    assert client.post(f"/api/runs/{run_id}/suppliers/SE/approve").status_code == 409


def test_upload_errors_and_summary_without_key(client, run, monkeypatch):
    for key in ("OPENAI_API_KEY", "MODEL"):
        monkeypatch.delenv(key, raising=False)
    path = "/api/datasets/IEK"
    assert client.post(path, files={"monthly_sales": ("sales.xlsx", b"")}).json()["code"] == "missing_file"
    files = {role: (f"{role}.xlsx", b"not xlsx") for role in ("monthly_sales", "monthly_stock", "in_transit", "moq")}
    response = client.post(path, files=files)
    assert response.status_code == 422 and response.json()["code"] == "not_xlsx"
    assert client.post(path, files={"other": ("x.xlsx", b"PKbad")}).json()["code"] == "unknown_role"
    assert client.post(f"/api/runs/{run['run_id']}/summary", params={"supplier": "SE"}).status_code == 503
    assert client.post("/api/runs", json={"lead_time_days": 0}).json()["code"] == "validation_error"
